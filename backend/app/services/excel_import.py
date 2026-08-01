"""Excel-импорт каталога: генерация шаблона, парсинг и валидация, создание позиций."""

import io
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.catalog import Category, Product
from app.models.enums import ProductType

# Порядок колонок шаблона.
COLUMNS = [
    "Наименование",
    "Категория",
    "SKU",
    "Ед.изм.",
    "Залоговая стоимость",
    "Суточная цена",
    "Показывать на сайте",
]

_TRUE = {"да", "yes", "1", "true", "истина", "+"}


@dataclass
class ParsedRow:
    row_number: int
    name: str
    category: str
    sku: str | None
    unit: str
    deposit_price: Decimal
    daily_price: Decimal
    show_on_site: bool
    # Причина, по которой строка не будет импортирована (None → строка валидна).
    # В предпросмотре такие строки показываем, но помечаем и в импорт не берём.
    error: str | None = None


@dataclass
class ParseResult:
    rows: list[ParsedRow] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def build_template() -> bytes:
    """Сформировать пустой .xlsx-шаблон с заголовками и примером строки."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Товары"
    ws.append(COLUMNS)
    ws.append(["Тарелка белая 25 см", "Тарелки", "PLT-25", "шт", 200, 10, "да"])
    for i, _ in enumerate(COLUMNS, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _to_decimal(value, row_no: int, col: str, errors: list[str]) -> Decimal:
    if value is None or value == "":
        return Decimal("0")
    try:
        return Decimal(str(value).replace(",", ".")).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        # Причина без префикса «Строка N» — его добавит вызывающий код.
        errors.append(f"некорректное число в «{col}»: {value!r}")
        return Decimal("0")


def parse(file_bytes: bytes) -> ParseResult:
    """Разобрать загруженный .xlsx в строки + список ошибок по строкам."""
    result = ParseResult()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - любая ошибка чтения файла
        result.errors.append(f"Не удалось прочитать файл: {exc}")
        return result

    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        result.errors.append("Файл пуст")
        return result

    for idx, raw in enumerate(rows, start=2):
        if raw is None or all(c is None or str(c).strip() == "" for c in raw):
            continue
        cells = list(raw) + [None] * (len(COLUMNS) - len(raw))
        name = str(cells[0]).strip() if cells[0] is not None else ""
        # Собираем ошибки строки, но саму строку всё равно показываем в предпросмотре
        # (помеченной). В импорт пойдут только строки без ошибок.
        row_errors: list[str] = []
        if not name:
            row_errors.append("пустое наименование")
        deposit_price = _to_decimal(cells[4], idx, "Залоговая стоимость", row_errors)
        daily_price = _to_decimal(cells[5], idx, "Суточная цена", row_errors)
        error = "; ".join(row_errors) if row_errors else None
        if error:
            result.errors.append(f"Строка {idx}: {error}")
        result.rows.append(
            ParsedRow(
                row_number=idx,
                name=name,
                category=(str(cells[1]).strip() if cells[1] else ""),
                sku=(str(cells[2]).strip() if cells[2] else None),
                unit=(str(cells[3]).strip() if cells[3] else "шт"),
                deposit_price=deposit_price,
                daily_price=daily_price,
                show_on_site=(str(cells[6]).strip().lower() in _TRUE if cells[6] else False),
                error=error,
            )
        )
    return result


def _get_or_create_category(db: Session, name: str) -> Category | None:
    if not name:
        return None
    cat = db.execute(select(Category).where(Category.name == name)).scalar_one_or_none()
    if cat is None:
        cat = Category(name=name)
        db.add(cat)
        db.flush()
    return cat


def create_products(db: Session, rows: list[ParsedRow], *, user_id: int | None = None) -> int:
    """Создать товары из разобранных строк. По SKU обновляет существующий товар."""
    created = 0
    for r in rows:
        existing = None
        if r.sku:
            existing = db.execute(select(Product).where(Product.sku == r.sku)).scalar_one_or_none()
        category = _get_or_create_category(db, r.category)
        if existing:
            existing.name = r.name
            existing.unit = r.unit
            existing.deposit_price = r.deposit_price
            existing.daily_price = r.daily_price
            existing.show_on_site = r.show_on_site
            if category:
                existing.category_id = category.id
        else:
            db.add(
                Product(
                    name=r.name,
                    sku=r.sku,
                    unit=r.unit,
                    type=ProductType.ITEM,
                    deposit_price=r.deposit_price,
                    daily_price=r.daily_price,
                    show_on_site=r.show_on_site,
                    category_id=category.id if category else None,
                    created_by_id=user_id,
                )
            )
            created += 1
    db.commit()
    return created
